import io
import unittest
from unittest.mock import patch, MagicMock
from bson import ObjectId
from fastapi.testclient import TestClient
import openpyxl

from app.main import app
from app.core.auth import get_current_user
from app.services.pdf_service import generate_pdf_report
from app.services.excel_service import generate_excel_report

client = TestClient(app)

MOCK_USER_ID = str(ObjectId())
MOCK_USER = {
    "user_id": MOCK_USER_ID,
    "email": "report_tester@example.com",
    "name": "Report Tester",
    "role": "researcher"
}


class TestReportsSuite(unittest.TestCase):

    def setUp(self):
        app.dependency_overrides[get_current_user] = lambda: MOCK_USER

    def tearDown(self):
        app.dependency_overrides.clear()

    # -------------------------------------------------------------
    # 1. Authentication & Authorization Enforcement
    # -------------------------------------------------------------
    def test_01_authentication_enforced(self):
        """Verify report endpoints return 401 when unauthenticated."""
        app.dependency_overrides.clear()

        endpoints = [
            "/api/reports/funding",
            "/api/reports/patents",
            "/api/reports/research-trends",
            "/api/reports/innovation",
            "/api/reports/commercialization"
        ]
        for ep in endpoints:
            res = client.get(ep)
            self.assertIn(res.status_code, (401, 403), f"Endpoint {ep} should require authentication")

        res_pdf = client.post("/api/reports/export/pdf", json={"report_type": "funding"})
        self.assertIn(res_pdf.status_code, (401, 403))

        res_excel = client.post("/api/reports/export/excel", json={"report_type": "funding"})
        self.assertIn(res_excel.status_code, (401, 403))

    # -------------------------------------------------------------
    # 2. Preview Endpoints for All Report Types
    # -------------------------------------------------------------
    def test_02_funding_report_preview(self):
        res = client.get("/api/reports/funding?limit=10")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["report_type"], "funding")
        self.assertIn("title", data)
        self.assertIn("summary", data)
        self.assertIn("metrics", data)
        self.assertIn("records", data)
        self.assertIsInstance(data["records"], list)

    def test_03_patent_report_preview(self):
        res = client.get("/api/reports/patents?limit=10")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["report_type"], "patents")
        self.assertIn("title", data)
        self.assertIn("metrics", data)
        self.assertIn("records", data)

    def test_04_research_trends_report_preview(self):
        res = client.get("/api/reports/research-trends?limit=10")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["report_type"], "research-trends")
        self.assertIn("metrics", data)
        self.assertIn("records", data)

    def test_05_innovation_report_preview(self):
        res = client.get("/api/reports/innovation?technology=Computer%20technology&limit=10")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["report_type"], "innovation")
        self.assertIn("metrics", data)
        self.assertIn("records", data)

    def test_06_commercialization_report_preview(self):
        res = client.get("/api/reports/commercialization?technology=Computer%20technology&limit=10")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["report_type"], "commercialization")
        self.assertIn("metrics", data)
        self.assertIn("records", data)

    # -------------------------------------------------------------
    # 3. PDF Export Verification
    # -------------------------------------------------------------
    def test_07_pdf_export_all_types(self):
        report_types = ["funding", "patents", "research-trends", "innovation", "commercialization"]
        for r_type in report_types:
            payload = {"report_type": r_type, "limit": 10}
            if r_type in ("innovation", "commercialization"):
                payload["technology"] = "Computer technology"

            res = client.post("/api/reports/export/pdf", json=payload)
            self.assertEqual(res.status_code, 200, f"PDF export failed for {r_type}")
            self.assertEqual(res.headers.get("content-type"), "application/pdf")
            self.assertIn(f"{r_type}_report.pdf", res.headers.get("content-disposition", ""))
            # Validate PDF signature
            self.assertTrue(res.content.startswith(b"%PDF-"), f"Output for {r_type} is not a valid PDF")
            self.assertGreater(len(res.content), 1000, f"PDF content for {r_type} too short")

    # -------------------------------------------------------------
    # 4. Excel Export Verification
    # -------------------------------------------------------------
    def test_08_excel_export_all_types(self):
        report_types = ["funding", "patents", "research-trends", "innovation", "commercialization"]
        for r_type in report_types:
            payload = {"report_type": r_type, "limit": 10}
            if r_type in ("innovation", "commercialization"):
                payload["technology"] = "Computer technology"

            res = client.post("/api/reports/export/excel", json=payload)
            self.assertEqual(res.status_code, 200, f"Excel export failed for {r_type}")
            self.assertEqual(
                res.headers.get("content-type"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            self.assertIn(f"{r_type}_report.xlsx", res.headers.get("content-disposition", ""))

            # Validate openpyxl can load the workbook
            wb = openpyxl.load_workbook(io.BytesIO(res.content))
            self.assertIn("Executive Summary", wb.sheetnames)
            self.assertGreater(len(wb.sheetnames), 1)

            ws_summary = wb["Executive Summary"]
            self.assertIsNotNone(ws_summary.cell(row=1, column=1).value)

    # -------------------------------------------------------------
    # 5. Invalid Report Types
    # -------------------------------------------------------------
    def test_09_invalid_report_type_returns_400(self):
        bad_payload = {"report_type": "nonexistent_domain", "limit": 10}
        res_pdf = client.post("/api/reports/export/pdf", json=bad_payload)
        self.assertEqual(res_pdf.status_code, 400)
        self.assertIn("Unsupported report type", res_pdf.json().get("detail", ""))

        res_excel = client.post("/api/reports/export/excel", json=bad_payload)
        self.assertEqual(res_excel.status_code, 400)
        self.assertIn("Unsupported report type", res_excel.json().get("detail", ""))

    # -------------------------------------------------------------
    # 6. Empty Dataset & Missing Optional Fields Handling
    # -------------------------------------------------------------
    def test_10_empty_dataset_pdf_and_excel(self):
        empty_data = {
            "title": "Empty Intelligence Report",
            "report_type": "funding",
            "generated_at": "2026-10-13T00:00:00Z",
            "user_name": "Test User",
            "summary": "No matching records found for given filter parameters.",
            "metrics": [],
            "records": []
        }

        pdf_bytes = generate_pdf_report(empty_data)
        self.assertTrue(pdf_bytes.startswith(b"%PDF-"))

        excel_bytes = generate_excel_report(empty_data)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Executive Summary", wb.sheetnames)

    def test_11_missing_optional_fields_resilience(self):
        ragged_data = {
            "title": None,
            "report_type": "patents",
            "generated_at": None,
            "user_name": None,
            "summary": None,
            "metrics": [
                {"label": "Test Metric", "value": None, "subtext": None}
            ],
            "records": [
                {"title": None, "organization": None, "award_amount": None},
                {"patent_number": "US999", "inventors": []}
            ]
        }

        pdf_bytes = generate_pdf_report(ragged_data)
        self.assertTrue(pdf_bytes.startswith(b"%PDF-"))

        excel_bytes = generate_excel_report(ragged_data)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Executive Summary", wb.sheetnames)

    # -------------------------------------------------------------
    # 7. Error Handling without Exposing Internal Secrets
    # -------------------------------------------------------------
    @patch("app.routes.reports.generate_pdf_report")
    def test_12_pdf_export_failure_returns_clean_500(self, mock_pdf):
        mock_pdf.side_effect = RuntimeError("Low-level generator exception")
        res = client.post("/api/reports/export/pdf", json={"report_type": "funding"})
        self.assertEqual(res.status_code, 500)
        self.assertIn("Failed to generate PDF report", res.json().get("detail", ""))

    @patch("app.routes.reports.generate_excel_report")
    def test_13_excel_export_failure_returns_clean_500(self, mock_excel):
        mock_excel.side_effect = RuntimeError("Low-level generator exception")
        res = client.post("/api/reports/export/excel", json={"report_type": "funding"})
        self.assertEqual(res.status_code, 500)
        self.assertIn("Failed to generate Excel report", res.json().get("detail", ""))


if __name__ == "__main__":
    unittest.main()
