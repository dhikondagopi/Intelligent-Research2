import io
from datetime import datetime
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_excel_report(report_data: Dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()
    
    # ---------------------------------------------------------
    # Styles
    # ---------------------------------------------------------
    font_title = Font(name="Calibri", size=16, bold=True, color="0F172A")
    font_section = Font(name="Calibri", size=13, bold=True, color="1E293B")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    font_muted = Font(name="Calibri", size=9, italic=True, color="64748B")

    fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # =========================================================
    # Sheet 1: Executive Summary
    # =========================================================
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.cell(row=1, column=1, value=report_data.get("title", "Platform Intelligence Report")).font = font_title
    ws_summary.cell(row=2, column=1, value=f"Report Type: {report_data.get('report_type', '').upper()}").font = Font(name="Calibri", size=11, bold=True, color="2563EB")
    ws_summary.cell(row=3, column=1, value=f"Generated At: {report_data.get('generated_at', '')} | User: {report_data.get('user_name', '')}").font = font_muted

    ws_summary.cell(row=5, column=1, value="Executive Summary").font = font_section
    ws_summary.cell(row=6, column=1, value=report_data.get("summary", "")).font = font_regular

    # KPI Table
    metrics = report_data.get("metrics", [])
    if metrics:
        ws_summary.cell(row=8, column=1, value="Key Performance Analytics").font = font_section
        
        ws_summary.cell(row=9, column=1, value="Metric Label").font = font_header
        ws_summary.cell(row=9, column=1).fill = fill_header
        ws_summary.cell(row=9, column=2, value="Metric Value").font = font_header
        ws_summary.cell(row=9, column=2).fill = fill_header
        ws_summary.cell(row=9, column=3, value="Description / Subtext").font = font_header
        ws_summary.cell(row=9, column=3).fill = fill_header

        row_idx = 10
        for m in metrics:
            ws_summary.cell(row=row_idx, column=1, value=m.get("label", "")).font = font_bold
            ws_summary.cell(row=row_idx, column=2, value=m.get("value", "")).font = font_bold
            ws_summary.cell(row=row_idx, column=3, value=m.get("subtext", "")).font = font_regular
            
            for col in range(1, 4):
                ws_summary.cell(row=row_idx, column=col).border = cell_border
                ws_summary.cell(row=row_idx, column=col).fill = fill_kpi
            row_idx += 1

    # Applied Filters Box
    ws_summary.cell(row=16, column=1, value="Applied Filter Parameters").font = font_section
    f_idx = 17
    for k, v in report_data.get("filters_applied", {}).items():
        if v:
            ws_summary.cell(row=f_idx, column=1, value=str(k)).font = font_bold
            ws_summary.cell(row=f_idx, column=2, value=str(v)).font = font_regular
            f_idx += 1

    ws_summary.cell(row=f_idx+1, column=1, value=report_data.get("methodology", "")).font = font_muted

    # Auto-fit columns for Summary sheet
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 18)

    # =========================================================
    # Sheet 2: Detailed Intelligence Data
    # =========================================================
    ws_data = wb.create_sheet(title="Detailed Records")
    ws_data.views.sheetView[0].showGridLines = True

    records = report_data.get("records", [])
    report_type = report_data.get("report_type", "")

    headers = []
    keys = []

    if report_type == "funding":
        headers = ["Project ID", "Project Title", "Organization", "Agency", "Funding Type", "Award Amount", "Fiscal Year", "Contact PI", "Source"]
        keys = ["project_id", "title", "organization", "agency", "funding_type", "award_amount", "fiscal_year", "contact_pi", "source"]
    elif report_type == "patents":
        headers = ["Patent Number", "Patent Title", "Assignee", "Technology Field", "Filing Date", "Citations", "Inventors", "Source"]
        keys = ["patent_number", "title", "assignee", "technology_field", "filing_date", "citations", "inventor_count", "source"]
    elif report_type == "research-trends":
        headers = ["Research ID", "Publication Title", "Primary Topic", "Publication Year", "Citations", "Source"]
        keys = ["research_id", "title", "primary_topic", "publication_year", "citations", "source"]
    elif report_type == "innovation":
        headers = ["Technology Field", "Innovation Score", "Research Novelty", "Patent Strength", "Maturity Score", "Maturity Stage", "Market Potential", "Funding Relevance"]
        keys = ["technology", "innovation_score", "research_novelty", "patent_strength", "technology_maturity", "maturity_stage", "market_potential", "funding_relevance"]
    else: # commercialization
        headers = ["Technology Field", "Patent Count", "Organization Count", "Inventor Count", "Productization Signal", "Licensing Signal", "Startup Signal", "Partnership Signal"]
        keys = ["technology", "patent_count", "organization_count", "inventor_count", "productization_readiness", "licensing_readiness", "startup_readiness", "partnership_readiness"]

    # Write Headers
    for c_idx, h in enumerate(headers, 1):
        cell = ws_data.cell(row=1, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # Write Data Rows
    for r_idx, rec in enumerate(records, 2):
        is_alt = (r_idx % 2 == 0)
        for c_idx, k in enumerate(keys, 1):
            val = rec.get(k, "")
            cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = cell_border
            if is_alt:
                cell.fill = fill_alt

            # Formatting numbers
            if k == "award_amount" and isinstance(val, (int, float)):
                cell.number_format = "$#,##0"
                cell.alignment = align_right
            elif k in ["citations", "patent_count", "organization_count", "inventor_count"] and isinstance(val, (int, float)):
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif k in ["innovation_score", "research_novelty", "patent_strength", "technology_maturity", "market_potential", "funding_relevance"] and isinstance(val, (int, float)):
                cell.number_format = "0.00"
                cell.alignment = align_right

    # Auto-fit columns for Data sheet
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    # Save to binary stream
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()
